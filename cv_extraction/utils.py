import io
import os
import re
import tempfile
import PyPDF2
import docx
import pythoncom
from docx import Document
from openpyxl import Workbook
from win32com.client import Dispatch

def extract_text(file):
    content_type = file.content_type

    if content_type == 'application/pdf':
        return extract_text_from_pdf(file)
    elif content_type.startswith('text/'):
        return extract_text_from_text(file)
    elif content_type.endswith('.docx'):
        return extract_text_from_docx(file)
    elif content_type.endswith('.doc'):
        return convert_doc_to_docx_and_extract_text(file)
    else:
        return None

def extract_text_from_pdf(pdf_file):
    try:
        text = ""
        if pdf_file.name.endswith('.pdf'):
            pdf_content = pdf_file.read()
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(pdf_content))
            for page in pdf_reader.pages:
                text += page.extract_text()
        return text
    except Exception as e:
        print("Error occurred while extracting text from PDF document:", e)
        return None

def extract_text_from_text(text_file):
    try:
        text_content = text_file.read().decode("utf-8")
        return text_content
    except Exception as e:
        print("Error occurred while extracting text from text file:", e)
        return None

def extract_text_from_docx(docx_file):
    try:

        docx_content = docx_file.read()
        docx_doc = docx.Document(io.BytesIO(docx_content))
        text = '\n'.join([paragraph.text for paragraph in docx_doc.paragraphs])
        return text
    except Exception as e:
        print("Error occurred while extracting text from DOCX document:", e)
        return None


def convert_doc_to_docx_and_extract_text(doc_file):
    pythoncom.CoInitialize() 
    temp_doc_path = None
    temp_docx_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as temp_doc:
            temp_doc.write(doc_file.read())
            temp_doc_path = temp_doc.name

        word_app = Dispatch('Word.Application')
        doc = word_app.Documents.Open(temp_doc_path)

        temp_docx_path = temp_doc_path.replace('.doc', '.docx')
        doc.SaveAs(temp_docx_path, FileFormat=16)
        doc.Close()
        del doc

        word_app.Quit()

        text = extract_text_from_docx(open(temp_docx_path, 'rb'))
        return text

    except Exception as e:
        print("Error occurred while converting .doc to .docx and extracting text:", e)
        return None
    finally:
        try:
            if temp_doc_path and os.path.exists(temp_doc_path):
                os.remove(temp_doc_path)
            if temp_docx_path and os.path.exists(temp_docx_path):
                os.remove(temp_docx_path)
        except Exception as e:
            print("Error occurred while cleaning up temporary files:", e)



def create_excel_file(data, output_path):
    try:
        workbook = Workbook()
        sheet = workbook.active
        headers = list(data[0].keys())
        sheet.append(headers)
        for row in data:
            values = [row[header] for header in headers]
            sheet.append(values)
        workbook.save(output_path)
    except Exception as e:
        print("Error occurred while creating Excel file:", e)

def extract_email(text):
    try:
        if isinstance(text, str) or isinstance(text, bytes):
            email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
            emails = re.findall(email_pattern, text)
            return emails
        else:
            return []
    except Exception as e:
        print("Error occurred while extracting email addresses:", e)
        return []

def extract_contact_number(text):
    try:
        if isinstance(text, str) or isinstance(text, bytes):
            contact_pattern = r'\b(?:\+\d{1,2}\s)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b'
            contact_numbers = re.findall(contact_pattern, text)
            return contact_numbers
        else:
            return []
    except Exception as e:
        print("Error occurred while extracting contact numbers:", e)
        return []
